# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

import math
import re
from typing import Iterable

import numpy as np
import pandas as pd

from openpyxl.utils.cell import range_boundaries


def is_empty(value) -> bool:
    """Determina si una celda debe considerarse vacía."""
    if value is None:
        return True

    if isinstance(value, str):
        return value.strip() == ""

    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def normalize_name(value: object, fallback: str) -> str:
    """
    Convierte un nombre de columna/tabla a un identificador estable.

    Ejemplo:
        'Fecha Nacimiento' -> 'fecha_nacimiento'
    """
    if is_empty(value):
        text = fallback
    else:
        text = str(value).strip()

    # Quitar acentos.
    replacements = str.maketrans(
        "áéíóúüñÁÉÍÓÚÜÑ",
        "aeiouunAEIOUUN",
    )

    text = text.translate(replacements)

    text = text.lower()

    text = re.sub(r"[^a-zA-Z0-9_]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")

    if not text:
        text = fallback

    if text[0].isdigit():
        text = f"col_{text}"

    return text


def unique_names(
    names: Iterable[object],
    prefix: str = "columna"
) -> list[str]:
    """Normaliza encabezados y evita duplicados."""

    result = []
    used = {}

    for index, name in enumerate(names, start=1):

        base = normalize_name(
            name,
            f"{prefix}_{index}"
        )

        count = used.get(base, 0) + 1
        used[base] = count

        result.append(
            base if count == 1 else f"{base}_{count}"
        )

    return result


def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Limpia encabezados, espacios y representa vacíos como NaN."""

    df = df.copy()

    df.columns = unique_names(df.columns)

    # Cadenas vacías/espacios -> NaN
    for column in df.columns:

        if pd.api.types.is_object_dtype(df[column]):

            df[column] = df[column].map(
                lambda x:
                    np.nan
                    if isinstance(x, str) and x.strip() == ""
                    else (
                        x.strip()
                        if isinstance(x, str)
                        else x
                    )
            )

    # Eliminar filas completamente vacías.
    df = df.dropna(
        axis=0,
        how="all"
    )

    # Las columnas completamente vacías NO se eliminan.
    # Forman parte de la estructura de la futura tabla SQL.

    return df.reset_index(drop=True)


def excel_table_range(
    ws,
    table_name: str
) -> tuple[int, int, int, int]:
    """Devuelve los límites de una tabla estructurada de Excel."""

    table = ws.tables[table_name]

    return range_boundaries(table.ref)


def get_structured_tables(ws):
    """Obtiene las tablas estructuradas de una hoja."""

    return list(ws.tables.values())


def dataframe_from_excel_table(
    ws,
    table
) -> pd.DataFrame:
    """
    Lee exactamente el rango de una tabla estructurada.

    La primera fila del rango se considera encabezado.
    """

    min_col, min_row, max_col, max_row = (
        range_boundaries(table.ref)
    )

    rows = []

    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        values_only=True,
    ):
        rows.append(list(row))

    if not rows:
        return pd.DataFrame()

    header = rows[0]
    data = rows[1:]

    df = pd.DataFrame(
        data,
        columns=header
    )

    return clean_dataframe(df)


def row_values(
    ws,
    row_number: int,
    max_col: int
) -> list:

    return [
        ws.cell(
            row=row_number,
            column=col
        ).value
        for col in range(1, max_col + 1)
    ]


def nonempty_positions(
    values: list
) -> list[int]:

    return [
        i
        for i, value in enumerate(values)
        if not is_empty(value)
    ]


def looks_like_header(
    ws,
    row_number: int,
    max_col: int,
    min_header_cells: int
) -> tuple[float, int]:
    """
    Puntúa una fila como posible encabezado.

    La heurística considera:

    - cantidad de celdas ocupadas;
    - cantidad de textos;
    - fórmulas;
    - existencia de datos debajo;
    - repetición de encabezados;
    - palabras que suelen aparecer en nombres de columnas.
    """

    values = row_values(
        ws,
        row_number,
        max_col
    )

    positions = nonempty_positions(values)

    if len(positions) < min_header_cells:
        return -math.inf, 0

    text_count = sum(
        isinstance(values[i], str)
        for i in positions
    )

    formula_count = sum(
        isinstance(values[i], str)
        and values[i].startswith("=")
        for i in positions
    )

    avg_len = (
        np.mean(
            [
                len(str(values[i]).strip())
                for i in positions
            ]
        )
        if positions
        else 0
    )

    score = 0.0

    text_ratio = (
        text_count /
        max(len(positions), 1)
    )

    score += text_ratio * 30.0

    score += min(
        len(positions),
        25
    ) * 1.5

    score -= formula_count * 5.0

    if avg_len > 80:
        score -= 4

    overlap_scores = []

    repeated_header_text = 0
    comparable_columns = 0

    for next_row in range(
        row_number + 1,
        min(
            row_number + 6,
            ws.max_row + 1
        )
    ):

        next_values = row_values(
            ws,
            next_row,
            max_col
        )

        overlap = sum(
            not is_empty(next_values[i])
            for i in positions
        )

        overlap_scores.append(
            overlap /
            max(len(positions), 1)
        )

        for i in positions:

            if (
                isinstance(values[i], str)
                and not values[i].startswith("=")
            ):

                comparable_columns += 1

                if (
                    isinstance(next_values[i], str)
                    and
                    next_values[i].strip().casefold()
                    ==
                    values[i].strip().casefold()
                ):
                    repeated_header_text += 1

    if overlap_scores:
        score += max(overlap_scores) * 12.0

    if comparable_columns:

        repetition_ratio = (
            repeated_header_text /
            comparable_columns
        )

        score -= repetition_ratio * 25.0

    header_keyword_pattern = re.compile(
        r"(id|codigo|cod|nombre|fecha|registro|"
        r"telefono|direccion|turno|puesto|maquina|"
        r"cargo|estado|tipo|valor|cantidad|empresa|"
        r"categoria|proceso|semana|dia|ruta|planta)",
        re.IGNORECASE,
    )

    keyword_hits = sum(
        bool(
            header_keyword_pattern.search(
                str(values[i])
            )
        )
        for i in positions
        if isinstance(values[i], str)
    )

    score += keyword_hits * 3.0

    return score, len(positions)


def detect_table_block(
    ws,
    max_scan_rows: int,
    max_scan_cols: int,
    min_header_cells: int
) -> tuple[int, int, int, int] | None:
    """
    Detecta una única tabla principal cuando la hoja
    no tiene una tabla estructurada de Excel.

    Retorna:

        (
            header_row,
            first_col,
            last_col,
            last_row
        )
    """

    scan_rows = min(
        ws.max_row,
        max_scan_rows
    )

    scan_cols = min(
        ws.max_column,
        max_scan_cols
    )

    candidates = []

    for row in range(
        1,
        scan_rows + 1
    ):

        score, count = looks_like_header(
            ws,
            row,
            scan_cols,
            min_header_cells
        )

        if score == -math.inf:
            continue

        positions = nonempty_positions(
            row_values(
                ws,
                row,
                scan_cols
            )
        )

        if not positions:
            continue

        first_col = min(positions) + 1
        last_col = max(positions) + 1

        candidates.append(
            (
                score,
                count,
                row,
                first_col,
                last_col
            )
        )

    if not candidates:
        return None

    candidates.sort(reverse=True)

    (
        _,
        header_count,
        header_row,
        first_col,
        last_col
    ) = candidates[0]

    if header_count < min_header_cells:
        return None

    last_row = header_row
    empty_streak = 0

    for row in range(
        header_row + 1,
        ws.max_row + 1
    ):

        values = [
            ws.cell(
                row=row,
                column=col
            ).value
            for col in range(
                first_col,
                last_col + 1
            )
        ]

        if all(
            is_empty(v)
            for v in values
        ):

            empty_streak += 1

            if empty_streak >= 3:
                break

            continue

        empty_streak = 0
        last_row = row

    if last_row <= header_row:
        return None

    return (
        header_row,
        first_col,
        last_col,
        last_row
    )


def dataframe_from_detected_block(
    ws,
    block
) -> pd.DataFrame:
    """Construye un DataFrame desde un bloque detectado."""

    (
        header_row,
        first_col,
        last_col,
        last_row
    ) = block

    rows = []

    for row in ws.iter_rows(
        min_row=header_row,
        max_row=last_row,
        min_col=first_col,
        max_col=last_col,
        values_only=True,
    ):
        rows.append(list(row))

    if not rows:
        return pd.DataFrame()

    header = rows[0]
    data = rows[1:]

    return clean_dataframe(
        pd.DataFrame(
            data,
            columns=header
        )
    )


def table_name(sheet_name: str) -> str:
    """Genera el nombre de tabla a partir del nombre de la hoja."""

    return normalize_name(
        sheet_name,
        "tabla"
    )


def process_sheet(
    ws,
    min_header_cells: int,
    max_scan_rows: int,
    max_scan_cols: int
) -> tuple[pd.DataFrame, str, str]:
    """
    Procesa una hoja.

    Prioridad:

    1. Tabla estructurada de Excel.
    2. Detección automática de bloque.
    """

    structured_tables = get_structured_tables(ws)

    if structured_tables:

        # Si hay varias tablas estructuradas,
        # se toma la de mayor superficie.
        table = max(
            structured_tables,
            key=lambda t: (
                range_boundaries(t.ref)[2]
                - range_boundaries(t.ref)[0]
                + 1
            )
            *
            (
                range_boundaries(t.ref)[3]
                - range_boundaries(t.ref)[1]
                + 1
            ),
        )

        df = dataframe_from_excel_table(
            ws,
            table
        )

        source = (
            f"tabla_excel:{table.name}"
        )

        return (
            df,
            source,
            table.ref
        )

    block = detect_table_block(
        ws,
        max_scan_rows,
        max_scan_cols,
        min_header_cells
    )

    if block is None:
        return (
            pd.DataFrame(),
            "no_detectada",
            ""
        )

    df = dataframe_from_detected_block(
        ws,
        block
    )

    source = "deteccion_automatica"

    ref = (
        f"{ws.cell(block[0], block[1]).coordinate}:"
        f"{ws.cell(block[3], block[2]).coordinate}"
    )

    return (
        df,
        source,
        ref
    )